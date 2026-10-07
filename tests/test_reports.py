import io
import openpyxl
from datetime import date, datetime, timezone
from app.models import AttendanceSession, AttendanceRecord, Student, ClassRoom, Subject, User
from app.services.report_service import (
    get_filtered_attendance_query, generate_attendance_csv,
    generate_attendance_xlsx, get_shortage_report
)
from app.utils import calculate_attendance_percentage, calculate_required_future_classes
from app.extensions import db


def test_attendance_percentage_formula():
    # 8 present, 2 late out of 10 eligible = 100%
    assert calculate_attendance_percentage(8, 2, 10) == 100.0
    # 7 present, 0 late out of 10 = 70%
    assert calculate_attendance_percentage(7, 0, 10) == 70.0
    # 0 sessions = 0%
    assert calculate_attendance_percentage(0, 0, 0) == 0.0


def test_future_classes_shortage_projection():
    # 5 attended out of 10 (50%). Target: 75%
    # (5 + x) / (10 + x) >= 0.75 => 5 + x >= 7.5 + 0.75x => 0.25x >= 2.5 => x >= 10
    needed = calculate_required_future_classes(5, 10, target_percentage=75.0)
    assert needed == 10


def test_csv_and_excel_export_generation(client):
    cls = ClassRoom.query.first()
    sub = Subject.query.first()
    teacher = User.query.filter_by(role='teacher').first()
    stu = Student.query.first()

    session = AttendanceSession(
        class_id=cls.id,
        subject_id=sub.id,
        teacher_id=teacher.id,
        date=date.today(),
        period=1,
        session_name="Export Test Session",
        status='ended'
    )
    db.session.add(session)
    db.session.flush()

    record = AttendanceRecord(
        session_id=session.id,
        student_id=stu.id,
        status='Present',
        method='face',
        confidence=0.92,
        note='Verified'
    )
    db.session.add(record)
    db.session.commit()

    query = get_filtered_attendance_query(class_id=cls.id)

    # 1. Test CSV export
    csv_content = generate_attendance_csv(query)
    assert "Record ID,Date,Period,Class,Section,Subject Code" in csv_content
    assert stu.roll_no in csv_content
    assert "Present" in csv_content

    # 2. Test Excel (.xlsx) export
    excel_stream = generate_attendance_xlsx(query, report_title="Test Report")
    assert isinstance(excel_stream, io.BytesIO)
    
    wb = openpyxl.load_workbook(excel_stream)
    assert "Attendance Records" in wb.sheetnames
    ws = wb["Attendance Records"]
    assert ws['A1'].value == "TEST REPORT"
